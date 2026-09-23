import csv
from datetime import timedelta
from decimal import Decimal
from django.contrib.auth.mixins import LoginRequiredMixin
from django.core.exceptions import PermissionDenied
from django.db.models import Sum, Count, Q
from django.http import HttpResponse
from django.shortcuts import render
from django.utils import timezone
from django.views.generic import View

from apps.accounts.permissions import DirectorOrAdminMixin
from apps.students.models import Eleve
from apps.payments.models import Paiement, TypeFrais
from apps.payments.services import get_situation_financiere


class ReportMixin(LoginRequiredMixin, View):
    def dispatch(self, request, *args, **kwargs):
        if not request.user.is_authenticated:
            return super().dispatch(request, *args, **kwargs)
        if request.user.role not in ["administrateur", "directeur"] and not request.user.is_superuser:
            raise PermissionDenied("Rapports réservés Direction/Administrateur")
        return super().dispatch(request, *args, **kwargs)


class ReportsDashboardView(ReportMixin, View):
    def get(self, request):
        today = timezone.now().date()
        # Recettes périodes
        qs = Paiement.objects.all()
        recettes_jour = qs.filter(date_paiement=today).aggregate(t=Sum("montant_paye"))["t"] or Decimal("0")
        recettes_semaine = qs.filter(date_paiement__gte=today - timedelta(days=7)).aggregate(t=Sum("montant_paye"))["t"] or Decimal("0")
        recettes_mois = qs.filter(date_paiement__year=today.year, date_paiement__month=today.month).aggregate(t=Sum("montant_paye"))["t"] or Decimal("0")
        recettes_annee = qs.filter(date_paiement__year=today.year).aggregate(t=Sum("montant_paye"))["t"] or Decimal("0")

        # Eleves à jour / débiteurs : sans dette paramétrée, on définit
        # à jour = ayant au moins un paiement sur l'année
        # débiteur = aucun paiement
        eleves = Eleve.objects.select_related("classe").all()
        a_jour_list = []
        debiteurs_list = []
        for e in eleves:
            sit = get_situation_financiere(e)
            if sit["total_paye"] > 0:
                a_jour_list.append((e, sit))
            else:
                debiteurs_list.append((e, sit))

        # Recettes par type (type_frais est CharField)
        recettes_par_type_qs = Paiement.objects.values("type_frais").annotate(total=Sum("montant_paye"), count=Count("id")).order_by("-total")
        # ajouter label
        label_map = dict(TypeFrais.choices)
        recettes_par_type = []
        for r in recettes_par_type_qs:
            r["type_frais_label"] = label_map.get(r["type_frais"], r["type_frais"])
            # compat ancien template qui attend type_frais__libelle
            r["type_frais__libelle"] = r["type_frais_label"]
            recettes_par_type.append(r)

        # Stats par classe
        from apps.classes.models import Classe
        stats_classe = []
        for cl in Classe.objects.all().order_by("niveau", "section"):
            eleves_cl = Eleve.objects.filter(classe=cl)
            total_du_cl = Decimal("0")
            total_paye_cl = Decimal("0")
            for e in eleves_cl:
                sit = get_situation_financiere(e)
                total_du_cl += sit["total_du"]
                total_paye_cl += sit["total_paye"]
            effectif = eleves_cl.count()
            # sans dette, taux = 100 si total_paye>0 sinon 0 (ou ratio à jour)
            if total_paye_cl > 0:
                taux = 100.0 if effectif else 0
                # plus fin : taux = % d'élèves ayant payé
                # calcul alternative si besoin
            else:
                taux = 0
            stats_classe.append({
                "classe": cl,
                "effectif": effectif,
                "total_du": total_du_cl,
                "total_paye": total_paye_cl,
                "solde": Decimal("0"),
                "taux": round(taux, 1),
            })

        ctx = {
            "recettes_jour": recettes_jour,
            "recettes_semaine": recettes_semaine,
            "recettes_mois": recettes_mois,
            "recettes_annee": recettes_annee,
            "a_jour_list": a_jour_list[:10],
            "debiteurs_list": debiteurs_list[:10],
            "recettes_par_type": recettes_par_type,
            "stats_classe": stats_classe,
            "total_eleves": eleves.count(),
            "a_jour_count": len(a_jour_list),
            "debiteurs_count": len(debiteurs_list),
        }
        return render(request, "reports/dashboard.html", ctx)


class ReportRecettesView(ReportMixin, View):
    def get(self, request):
        period = request.GET.get("period", "jour")  # jour, semaine, mois, annee
        today = timezone.now().date()
        qs = Paiement.objects.select_related("eleve", "agent", "recu_associe").order_by("-date_paiement")
        if period == "jour":
            qs = qs.filter(date_paiement=today)
        elif period == "semaine":
            qs = qs.filter(date_paiement__gte=today - timedelta(days=7))
        elif period == "mois":
            qs = qs.filter(date_paiement__year=today.year, date_paiement__month=today.month)
        elif period == "annee":
            qs = qs.filter(date_paiement__year=today.year)
        total = qs.aggregate(t=Sum("montant_paye"))["t"] or Decimal("0")
        from django.core.paginator import Paginator
        paginator = Paginator(qs, 12)
        page_number = request.GET.get("page")
        page_obj = paginator.get_page(page_number)
        ctx = {"paiements": page_obj.object_list, "page_obj": page_obj, "paginator": paginator, "is_paginated": page_obj.has_other_pages(), "total": total, "period": period}
        if request.headers.get("HX-Request") == "true" or request.META.get("HTTP_HX_REQUEST") == "true":
            return render(request, "reports/_recettes_card.html", ctx)
        return render(request, "reports/recettes.html", ctx)


class ReportDebiteursView(ReportMixin, View):
    def get(self, request):
        eleves = Eleve.objects.select_related("classe").all()
        debiteurs = []
        for e in eleves:
            sit = get_situation_financiere(e)
            if sit["total_paye"] == 0:
                debiteurs.append(sit)
        # tri par nom
        debiteurs.sort(key=lambda x: x["eleve"].nom)
        from django.core.paginator import Paginator
        paginator = Paginator(debiteurs, 12)
        page_number = request.GET.get("page")
        page_obj = paginator.get_page(page_number)
        return render(request, "reports/debiteurs.html", {"debiteurs": page_obj.object_list, "page_obj": page_obj, "paginator": paginator, "is_paginated": page_obj.has_other_pages()})


class ReportAjourView(ReportMixin, View):
    def get(self, request):
        eleves = Eleve.objects.select_related("classe").all()
        a_jour = []
        for e in eleves:
            sit = get_situation_financiere(e)
            if sit["total_paye"] > 0:
                a_jour.append(sit)
        from django.core.paginator import Paginator
        paginator = Paginator(a_jour, 12)
        page_number = request.GET.get("page")
        page_obj = paginator.get_page(page_number)
        return render(request, "reports/a_jour.html", {"a_jour": page_obj.object_list, "page_obj": page_obj, "paginator": paginator, "is_paginated": page_obj.has_other_pages()})


class ReportStatsClasseView(ReportMixin, View):
    def get(self, request):
        from apps.classes.models import Classe
        from django.core.paginator import Paginator
        stats = []
        for cl in Classe.objects.all().order_by("niveau", "section"):
            eleves_cl = Eleve.objects.filter(classe=cl)
            total_du = Decimal("0")
            total_paye = Decimal("0")
            a_jour = 0
            debiteurs = 0
            for e in eleves_cl:
                sit = get_situation_financiere(e)
                total_du += sit["total_du"]
                total_paye += sit["total_paye"]
                if sit["total_paye"] > 0:
                    a_jour += 1
                else:
                    debiteurs += 1
            taux = 100.0 if total_paye > 0 else 0
            stats.append({
                "classe": cl,
                "effectif": eleves_cl.count(),
                "total_du": total_du,
                "total_paye": total_paye,
                "a_jour": a_jour,
                "debiteurs": debiteurs,
                "taux": round(taux, 1),
            })
        paginator = Paginator(stats, 12)
        page_number = request.GET.get("page")
        page_obj = paginator.get_page(page_number)
        return render(request, "reports/stats_classe.html", {"stats": page_obj.object_list, "page_obj": page_obj, "paginator": paginator, "is_paginated": page_obj.has_other_pages()})


class ExportMixin:
    def log_export(self, request, name):
        from apps.audit.models import AuditLog
        AuditLog.objects.create(
            user=request.user,
            action="EXPORT",
            model_name="Report",
            object_repr=name,
            ip_address=request.META.get("REMOTE_ADDR"),
        )


class ExportExcelView(ReportMixin, ExportMixin, View):
    def get(self, request):
        report = request.GET.get("report", "debiteurs")
        self.log_export(request, f"export_excel_{report}")
        try:
            import openpyxl
            from openpyxl.styles import Font
        except ImportError:
            return HttpResponse("openpyxl non installé. Ajoutez openpyxl au requirements.", status=500, content_type="text/plain")

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = report

        label_map = dict(TypeFrais.choices)

        if report == "debiteurs":
            ws.append(["Matricule", "Nom complet", "Classe", "Total payé"])
            for cell in ws[1]:
                cell.font = Font(bold=True)
            for e in Eleve.objects.select_related("classe").all():
                sit = get_situation_financiere(e)
                if sit["total_paye"] == 0:
                    ws.append([e.matricule, e.nom_complet, e.classe.nom, float(sit["total_paye"])])
        elif report == "a_jour":
            ws.append(["Matricule", "Nom complet", "Classe", "Total payé"])
            for cell in ws[1]:
                cell.font = Font(bold=True)
            for e in Eleve.objects.select_related("classe").all():
                sit = get_situation_financiere(e)
                if sit["total_paye"] > 0:
                    ws.append([e.matricule, e.nom_complet, e.classe.nom, float(sit["total_paye"])])
        elif report == "recettes":
            ws.append(["Date", "Reçu", "Élève", "Type", "Montant", "Devise", "Agent"])
            for cell in ws[1]:
                cell.font = Font(bold=True)
            period = request.GET.get("period")
            qs = Paiement.objects.select_related("eleve", "agent", "recu_associe").order_by("-date_paiement")
            if period == "jour":
                qs = qs.filter(date_paiement=timezone.now().date())
            elif period == "semaine":
                qs = qs.filter(date_paiement__gte=timezone.now().date() - timedelta(days=7))
            elif period == "mois":
                qs = qs.filter(date_paiement__year=timezone.now().date().year, date_paiement__month=timezone.now().date().month)
            elif period == "annee":
                qs = qs.filter(date_paiement__year=timezone.now().date().year)
            for p in qs[:2000]:
                ws.append([p.date_paiement.isoformat(), getattr(p.recu_associe, "numero", "—"), p.eleve.matricule, str(label_map.get(p.type_frais, p.type_frais)), float(p.montant_paye), p.devise, p.agent.login])
        elif report == "stats_classe":
            ws.append(["Classe", "Effectif", "Total payé", "Taux %"])
            for cell in ws[1]:
                cell.font = Font(bold=True)
            from apps.classes.models import Classe
            for cl in Classe.objects.all().order_by("niveau", "section"):
                eleves_cl = Eleve.objects.filter(classe=cl)
                total_paye = sum((get_situation_financiere(e)["total_paye"] for e in eleves_cl), Decimal("0"))
                taux = 100.0 if total_paye > 0 else 0
                ws.append([cl.nom, eleves_cl.count(), float(total_paye), round(taux,1)])
        else:
            ws.append(["Matricule", "Nom complet", "Classe", "Total payé"])
            for cell in ws[1]:
                cell.font = Font(bold=True)
            for e in Eleve.objects.select_related("classe").all()[:100]:
                sit = get_situation_financiere(e)
                ws.append([e.matricule, e.nom_complet, e.classe.nom, float(sit["total_paye"])])

        from io import BytesIO
        buf = BytesIO()
        wb.save(buf)
        buf.seek(0)
        response = HttpResponse(buf.getvalue(), content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
        response["Content-Disposition"] = f'attachment; filename="meriba_{report}.xlsx"'
        return response


class ExportPDFView(ReportMixin, ExportMixin, View):
    def get(self, request):
        report = request.GET.get("report", "debiteurs")
        self.log_export(request, f"export_pdf_{report}")
        try:
            from reportlab.lib.pagesizes import A4
            from reportlab.lib.units import mm
            from reportlab.lib import colors
            from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer
            from reportlab.lib.styles import getSampleStyleSheet
            from io import BytesIO
        except ImportError:
            return HttpResponse("reportlab non installé.", status=500, content_type="text/plain")

        from io import BytesIO
        buffer = BytesIO()
        from reportlab.lib.pagesizes import A4
        from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer
        from reportlab.lib.styles import getSampleStyleSheet
        from reportlab.lib import colors
        doc = SimpleDocTemplate(buffer, pagesize=A4, topMargin=20, bottomMargin=20)
        styles = getSampleStyleSheet()
        story = []
        story.append(Paragraph("Complexe Scolaire Meriba — Rapport " + report.capitalize(), styles["Title"]))
        story.append(Spacer(1, 12))
        story.append(Paragraph(f"Généré le {timezone.now().strftime('%d/%m/%Y %H:%M')} par {request.user.login}", styles["Normal"]))
        story.append(Spacer(1, 12))

        label_map = dict(TypeFrais.choices)

        if report == "debiteurs":
            data = [["Matricule", "Nom", "Classe", "Total paye"]]
            for e in Eleve.objects.select_related("classe").all():
                sit = get_situation_financiere(e)
                if sit["total_paye"] == 0:
                    data.append([e.matricule, e.nom_complet, e.classe.nom, str(sit["total_paye"])])
            if len(data) == 1:
                data.append(["—", "Aucun débiteur", "—", "0"])
        elif report == "a_jour":
            data = [["Matricule", "Nom", "Classe", "Total paye"]]
            for e in Eleve.objects.select_related("classe").all():
                sit = get_situation_financiere(e)
                if sit["total_paye"] > 0:
                    data.append([e.matricule, e.nom_complet, e.classe.nom, str(sit["total_paye"])])
            if len(data) == 1:
                data.append(["—", "Aucun élève à jour", "—", "0"])
        elif report == "recettes":
            data = [["Date", "Reçu", "Élève", "Type", "Montant", "Devise"]]
            qs = Paiement.objects.select_related("eleve", "recu_associe").order_by("-date_paiement")
            period = request.GET.get("period")
            if period == "jour":
                qs = qs.filter(date_paiement=timezone.now().date())
            elif period == "semaine":
                qs = qs.filter(date_paiement__gte=timezone.now().date() - timedelta(days=7))
            elif period == "mois":
                qs = qs.filter(date_paiement__year=timezone.now().date().year, date_paiement__month=timezone.now().date().month)
            elif period == "annee":
                qs = qs.filter(date_paiement__year=timezone.now().date().year)
            for p in qs[:2000]:
                data.append([p.date_paiement.isoformat(), getattr(p.recu_associe, "numero", "—"), p.eleve.matricule, str(label_map.get(p.type_frais, p.type_frais)), str(p.montant_paye), p.devise])
            if len(data) == 1:
                data.append(["—", "—", "—", "—", "0", "—"])
        elif report == "stats_classe":
            data = [["Classe", "Effectif", "Total paye", "Taux %"]]
            from apps.classes.models import Classe
            for cl in Classe.objects.all().order_by("niveau", "section"):
                eleves_cl = Eleve.objects.filter(classe=cl)
                total_paye = sum((get_situation_financiere(e)["total_paye"] for e in eleves_cl), Decimal("0"))
                data.append([cl.nom, str(eleves_cl.count()), str(total_paye), str(round(100.0 if total_paye>0 else 0,1))])
        else:
            data = [["Matricule", "Nom", "Classe", "Total paye"]]
            for e in Eleve.objects.select_related("classe").all()[:100]:
                sit = get_situation_financiere(e)
                data.append([e.matricule, e.nom_complet, e.classe.nom, str(sit["total_paye"])])

        t = Table(data, repeatRows=1)
        t.setStyle(TableStyle([
            ("BACKGROUND", (0,0), (-1,0), colors.HexColor("#f6f8fa")),
            ("GRID", (0,0), (-1,-1), 0.5, colors.HexColor("#d0d7de")),
            ("FONTNAME", (0,0), (-1,0), "Helvetica-Bold"),
            ("FONTSIZE", (0,0), (-1,-1), 8),
            ("ALIGN", (-1,1), (-1,-1), "RIGHT"),
            ("LEFTPADDING", (0,0), (-1,-1), 4),
            ("RIGHTPADDING", (0,0), (-1,-1), 4),
        ]))
        story.append(t)
        doc.build(story)
        buffer.seek(0)
        response = HttpResponse(buffer.getvalue(), content_type="application/pdf")
        response["Content-Disposition"] = f'attachment; filename="meriba_{report}.pdf"'
        return response
